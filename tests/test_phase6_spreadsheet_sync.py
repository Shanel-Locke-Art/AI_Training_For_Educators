#!/usr/bin/env python3
"""Guard the current canonical asset documentation and Office trackers."""

from pathlib import Path
from zipfile import ZipFile
import json

from openpyxl import load_workbook
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/asset-management"
MANIFEST = json.loads((ROOT / "assets/asset-manifest.json").read_text(encoding="utf-8"))

assert MANIFEST["version"] == 150
assert MANIFEST["runtime_images"]["mo-river-otter.png"] == "assets/images/ui/mo-river-otter.png"

for name, sheets in {
    "PromptCraft_Production_Overview_Simplified.xlsx": ["Overview", "Current Work", "File Guide"],
    "PromptCraft_Visual_Asset_Tracker_Simplified.xlsx": ["Summary", "Visual Inventory"],
    "PromptCraft_Voice_Recording_Tracker.xlsx": ["Start Here", "Audio Inventory", "Recording Plan", "Retired Names"],
}.items():
    path = DOCS / name
    with ZipFile(path) as archive:
        assert archive.testzip() is None
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == sheets
    text = " ".join(str(cell.value or "") for sheet in workbook for row in sheet for cell in row)
    for token in ("PROMPTCRAFT_V429", "Patch 575", "v150"):
        assert token in text, f"{name} missing {token}"

visual = load_workbook(DOCS / "PromptCraft_Visual_Asset_Tracker_Simplified.xlsx", data_only=False)
paths = {row[3].value for row in visual["Visual Inventory"].iter_rows(min_row=5) if row[3].value}
expected = {
    p.relative_to(ROOT).as_posix()
    for p in (ROOT / "assets").rglob("*")
    if p.is_file() and p.suffix.lower() in {".png", ".jpg", ".jpeg", ".svg", ".webp", ".gif"}
}
assert paths == expected

for name in ("Professor_Pixel_Recording_Script.docx", "Jordan_Recording_Script.docx", "Eli_Recording_Script.docx"):
    path = DOCS / "Recording Scripts" / name
    with ZipFile(path) as archive:
        assert archive.testzip() is None
    document = Document(path)
    text = " ".join(
        [p.text for p in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    assert "Patch 575" in text
    assert "Before recording" in text

assert (DOCS / "ASSET_SYSTEM.md").exists()
print("Current asset documentation and Office tracker contract passed")

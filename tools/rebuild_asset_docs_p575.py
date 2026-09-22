from pathlib import Path
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path.cwd(); DOCS=ROOT/'docs/asset-management'; ASSETS=ROOT/'assets'; MAN=ASSETS/'asset-manifest.json'
m=json.loads(MAN.read_text(encoding='utf-8')); m['version']=150
m['runtime_images']['mo-river-otter.png']='assets/images/ui/mo-river-otter.png'
MAN.write_text(json.dumps(m,indent=2)+"\n",encoding='utf-8')

all_asset_files=sorted(p for p in ASSETS.rglob('*') if p.is_file() and p.name not in ('desktop.ini','asset-manifest.json') and not p.name.lower().endswith('.md'))
images=[p for p in all_asset_files if p.suffix.lower() in {'.png','.jpg','.jpeg','.svg','.webp','.gif'}]
audio=[p for p in all_asset_files if p.suffix.lower() in {'.mp3','.wav','.ogg','.m4a'}]
runtime=set(m['runtime_images'].values())|set(m['runtime_audio'].values())
planned=set(m['planned_runtime_images'].values())|set(m['planned_audio'])
refs=set(m['development_reference_images'])

def rel(p): return p.relative_to(ROOT).as_posix()
def category(p):
 s=rel(p)
 if '/backgrounds/' in s:return 'Background'
 if '/characters/professor-pixel/' in s:return 'Character - Professor Pixel'
 if '/characters/students/' in s:return 'Character - Student'
 if '/canvas/' in s:return 'Canvas evidence'
 if '/scenes/' in s:return 'Scene art'
 if '/ui/' in s:return 'UI and brand'
 if '/brand/' in s:return 'Brand'
 return 'Other visual'
def use_for(p):
 s=rel(p)
 if 'scenario-01-content-avalanche/canvas' in s:return 'Legacy S1 Canvas evidence set retained for reference and dormant modules'
 if 'maya/' in s:return 'Maya portrait used by current Scenario 1 Start With the Learning'
 if 'professor-pixel' in s:return 'Professor Pixel portrait'
 if 'gfc/s1-science-wing' in s:return 'Current Scenario 1 background'
 if '/scenes/' in s:return 'Scenario scene or completion art'
 return 'Shared application asset'
def status(p):
 s=rel(p)
 if s in runtime:return 'Runtime'
 if s in planned:return 'Planned'
 if s in refs or '/references/' in s:return 'Reference'
 return 'Present, not registered'

def setup(ws,title,subtitle,widths):
 ws.sheet_view.showGridLines=False; ws.freeze_panes='A5'
 ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(widths)); c=ws.cell(1,1,title); c.font=Font(name='Aptos Display',size=20,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='0D2A55');c.alignment=Alignment(vertical='center');ws.row_dimensions[1].height=34
 ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(widths)); c=ws.cell(2,1,subtitle);c.font=Font(name='Aptos',size=10,color='334D60');c.fill=PatternFill('solid',fgColor='EAF2F7');c.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=32
 for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w

def header(ws,row,labels):
 for i,v in enumerate(labels,1):
  c=ws.cell(row,i,v);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1672A5');c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[row].height=28

def body_style(ws,start,end,cols):
 side=Side(style='thin',color='D9D9D9')
 for row in ws.iter_rows(min_row=start,max_row=end,min_col=1,max_col=cols):
  for c in row:
   c.border=Border(left=side,right=side,top=side,bottom=side);c.alignment=Alignment(vertical='center',wrap_text=True)
  if row[0].row%2==0:
   for c in row:c.fill=PatternFill('solid',fgColor='F4F8FB')

def save_visual_tracker():
 wb=Workbook(); s=wb.active;s.title='Summary';setup(s,'PromptCraft Visual Asset Tracker','Current asset inventory generated from the filesystem and asset manifest v150. Application baseline: PROMPTCRAFT_V429, Patch 575.',[26,16,16,16,18,18])
 header(s,4,['Measure','Count','Meaning','','',''])
 rows=[('Visual files present',len(images),'All image files under assets/images and assets/ui'),('Registered runtime visuals',sum(rel(p) in runtime for p in images),'Loaded or available through the runtime registry'),('Planned visuals',sum(rel(p) in planned for p in images),'Reserved for future scenario use'),('Reference visuals',sum(rel(p) in refs or '/references/' in rel(p) for p in images),'Development reference only'),('Unregistered visuals',sum(status(p)=='Present, not registered' for p in images),'Needs an explicit keep, register, archive, or remove decision')]
 for r,row in enumerate(rows,5):
  for c,v in enumerate(row,1):s.cell(r,c,v)
 body_style(s,5,9,3)
 s['A12']='Status rules';s['A12'].font=Font(bold=True,size=13,color='0D2A55')
 s['A13']='Runtime';s['B13']='The current application registry or runtime audio map names the file.'
 s['A14']='Planned';s['B14']='The manifest reserves the file for future scenario work.'
 s['A15']='Reference';s['B15']='Concept, source, or comparison material. Do not ship as a production portrait without review.'
 s['A16']='Present, not registered';s['B16']='The file exists but has no current manifest role. Resolve deliberately.'
 inv=wb.create_sheet('Visual Inventory');setup(inv,'Visual Inventory','One row per current visual file. Filter by status, category, or scenario use.',[34,24,36,64,22,18,42])
 labels=['File','Category','Current use','Path','Status','Needs decision','Notes'];header(inv,4,labels)
 for i,p in enumerate(images,5):
  st=status(p); vals=[p.name,category(p),use_for(p),rel(p),st,'Yes' if st=='Present, not registered' else 'No','Manifest v150 classification' if st!='Present, not registered' else 'Assign a manifest role or archive/remove.']
  for j,v in enumerate(vals,1):inv.cell(i,j,v)
 body_style(inv,5,4+len(images),len(labels));inv.auto_filter.ref=f'A4:G{4+len(images)}';inv.freeze_panes='A5'
 dv=DataValidation(type='list',formula1='"Runtime,Planned,Reference,Present, not registered,Archived"'); inv.add_data_validation(dv);dv.add(f'E5:E{4+len(images)}')
 wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto'
 wb.save(DOCS/'PromptCraft_Visual_Asset_Tracker_Simplified.xlsx')

def save_voice_tracker():
 wb=Workbook();s=wb.active;s.title='Start Here';setup(s,'PromptCraft Voice and Audio Tracker','Current audio inventory and recording decisions. Baseline: PROMPTCRAFT_V429, Patch 575, asset manifest v150.',[28,18,52,24,24,24])
 header(s,4,['Measure','Count','Current interpretation','','',''])
 vals=[('Audio files present',len(audio),'Physical audio files in the repository'),('Runtime audio',sum(rel(p) in runtime for p in audio),'Currently registered playback files'),('Planned recording paths',len(m['planned_audio']),'Paths reserved in the manifest; a missing file is expected'),('Current S1 voice status',0,'Scenario 1 Start With the Learning uses text dialogue and Maya panels; no new S1 voice files are approved for recording'),('Retired names',len(m['retired_audio_names']),'Names blocked from reuse')]
 for i,row in enumerate(vals,5):
  for j,v in enumerate(row,1):s.cell(i,j,v)
 body_style(s,5,9,3)
 inv=wb.create_sheet('Audio Inventory');setup(inv,'Audio Inventory','Actual files only. Runtime status comes from asset manifest v150.',[34,24,34,68,20,42]);header(inv,4,['File','Speaker or type','Use','Path','Status','Notes'])
 for i,p in enumerate(audio,5):
  rp=rel(p); who='Professor Pixel' if 'professor-pixel' in rp else ('Student voice' if '/students/' in rp else 'Music')
  st='Runtime' if rp in runtime else ('Planned file present' if rp in planned else 'Present, not registered')
  for j,v in enumerate([p.name,who,'Application audio',rp,st,'Verify wording before enabling voice playback.'],1):inv.cell(i,j,v)
 body_style(inv,5,4+len(audio),6);inv.auto_filter.ref=f'A4:F{4+len(audio)}'
 plan=wb.create_sheet('Recording Plan');setup(plan,'Recording Plan','Do not record from older scripts. Approve current dialogue, assign stable IDs, then create one file per approved line.',[22,28,30,60,22,48]);header(plan,4,['Priority','Speaker','Scenario','Action','Status','Decision rule'])
 planrows=[('High','Professor Pixel and Maya','S1 Start With the Learning','Export a fresh line list from current Scenario 1 source after wording lock.','Hold','The current scenario changed substantially; old Content Avalanche and Eli lines are retired.'),('High','Professor Pixel and Jordan','S3 Confident Student Problem','Reconcile the legacy scenario-02 paths with current displayed numbering.','Needs review','Keep path compatibility until code and receiver migration are approved.'),('Medium','Professor Pixel and future students','S4-S8','Record only after each scenario dialogue is locked.','Hold','Do not create speculative audio from draft dialogue.')]
 for i,row in enumerate(planrows,5):
  for j,v in enumerate(row,1):plan.cell(i,j,v)
 body_style(plan,5,7,6)
 ret=wb.create_sheet('Retired Names');setup(ret,'Retired Audio Names','These filenames must not be reintroduced without a documented migration.',[34,72]);header(ret,4,['Filename','Reason'])
 for i,n in enumerate(m['retired_audio_names'],5):ret.cell(i,1,n);ret.cell(i,2,'Retired by asset manifest v150; preserve only as historical reference.')
 body_style(ret,5,4+len(m['retired_audio_names']),2)
 wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto';wb.save(DOCS/'PromptCraft_Voice_Recording_Tracker.xlsx')

def save_overview():
 wb=Workbook();s=wb.active;s.title='Overview';setup(s,'PromptCraft Asset Production Overview','Current production view generated from asset manifest v150 and the repository asset tree. Baseline: PROMPTCRAFT_V429, Patch 575.',[30,18,56,26,26,26])
 header(s,4,['Area','Count','What the count represents','','',''])
 rows=[('Visual files',len(images),'All current visual files'),('Runtime visuals',sum(rel(p) in runtime for p in images),'Visuals registered for runtime use'),('Audio files',len(audio),'Physical audio files'),('Runtime audio',sum(rel(p) in runtime for p in audio),'Audio registered for playback'),('Open visual classifications',sum(status(p)=='Present, not registered' for p in images),'Files requiring an explicit role decision'),('Planned recording paths',len(m['planned_audio']),'Reserved future filenames, not completed recordings')]
 for i,row in enumerate(rows,5):
  for j,v in enumerate(row,1):s.cell(i,j,v)
 body_style(s,5,10,3)
 q=wb.create_sheet('Current Work');setup(q,'Current Asset Work','Resolve classification and recording decisions before creating more files.',[16,28,58,20,54]);header(q,4,['Priority','Area','Required action','Status','Owner check'])
 work=[('High','Scenario 1 voice','Replace the retired Content Avalanche and Eli recording plan with an approved Professor Pixel and Maya line export from current source.','Hold','Instructional owner approves final wording before recording.'),('High','Unregistered visuals',f'Review {sum(status(p)=="Present, not registered" for p in images)} files that exist without a manifest role.','Open','Assign Runtime, Planned, Reference, Archive, or Remove.'),('Medium','Legacy Canvas evidence','Confirm whether the Content Avalanche screenshots remain needed for dormant modules and research evidence.','Review','Do not relabel them as current Start With the Learning screens.'),('Medium','Future scenario audio','Keep draft paths on hold until S3-S8 dialogue and numbering are locked.','Hold','Preserve compatibility paths until migration is explicit.')]
 for i,row in enumerate(work,5):
  for j,v in enumerate(row,1):q.cell(i,j,v)
 body_style(q,5,8,5)
 f=wb.create_sheet('File Guide');setup(f,'Asset Documentation File Guide','Use each canonical file for one purpose. Do not create version-number duplicates.',[46,64,52]);header(f,4,['File','Use it for','Authority'])
 entries=[('PromptCraft_Production_Overview_Simplified.xlsx','Production totals, open work, and next decisions','Summary derived from the other trackers'),('PromptCraft_Visual_Asset_Tracker_Simplified.xlsx','Every current visual file and its manifest status','Filesystem plus assets/asset-manifest.json'),('PromptCraft_Voice_Recording_Tracker.xlsx','Actual audio files, recording holds, and retired names','Filesystem plus assets/asset-manifest.json'),('ASSET_SYSTEM.md','Naming, lifecycle, approval, and update workflow','Canonical written asset policy')]
 for i,row in enumerate(entries,5):
  for j,v in enumerate(row,1):f.cell(i,j,v)
 body_style(f,5,8,3)
 wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto';wb.save(DOCS/'PromptCraft_Production_Overview_Simplified.xlsx')

save_visual_tracker();save_voice_tracker();save_overview()

# Recording guides: process documents that deliberately avoid stale dialogue duplication.
def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),color);tcPr.append(shd)
def make_guide(path,title,speaker,current_role,decision):
 d=Document();sec=d.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
 st=d.styles['Normal'];st.font.name='Aptos';st.font.size=Pt(10.5);st.font.color.rgb=RGBColor(32,48,60)
 t=d.add_paragraph(title,style='Title');t.alignment=WD_ALIGN_PARAGRAPH.CENTER;t.runs[0].font.color.rgb=RGBColor(0,0,0)
 p=d.add_paragraph();p.add_run('Purpose. ').bold=True;p.add_run('Use this guide to decide whether a recording session is ready. The voice tracker and current application source control the approved line list; this document does not copy dialogue that may become stale.')
 d.add_heading('Current production status',level=1)
 table=d.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.style='Table Grid';
 trPr=table.rows[0]._tr.get_or_add_trPr();tbl_header=OxmlElement('w:tblHeader');tbl_header.set(qn('w:val'),'true');trPr.append(tbl_header)
 for c,txt in zip(table.rows[0].cells,['Item','Current state']):c.text=txt;shade(c,'0D2A55');c.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255);c.paragraphs[0].runs[0].bold=True
 for a,b in [('Speaker',speaker),('Current role',current_role),('Recording decision',decision),('Technical baseline','PROMPTCRAFT_V429, Patch 575, asset manifest v150')]:
  cells=table.add_row().cells;cells[0].text=a;cells[1].text=b
  for c in cells:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 d.add_heading('Before recording',level=1)
 for text in ['Open PromptCraft_Voice_Recording_Tracker.xlsx and confirm the line is marked Approved to Record.','Compare the exported line with the current browser source. Stop if wording, speaker, or scenario numbering differs.','Use one stable line ID and one audio file per approved line. Preserve legacy path names only when the tracker identifies a compatibility requirement.','Record dialogue only. Do not read labels, cues, IDs, or production notes.']:
  d.add_paragraph(text,style='List Bullet')
 d.add_heading('Performance direction',level=1)
 d.add_paragraph('Use a natural teaching or student voice. Keep the pacing conversational, avoid announcer delivery, and leave brief clean silence at the start and end for editing.')
 d.add_heading('File delivery checklist',level=1)
 for text in ['Filename matches the approved tracker row','No clipped first or last word','No background music or room noise','Expression and pacing match the approved cue','File opens and plays before it is added to the repository']:
  d.add_paragraph('☐ '+text)
 d.save(path)
make_guide(DOCS/'Recording Scripts/Professor_Pixel_Recording_Script.docx','Professor Pixel Recording Guide','Professor Pixel','Guide and instructional design coach across PromptCraft.','Hold new Scenario 1 recording until the current Start With the Learning dialogue is exported and approved.')
make_guide(DOCS/'Recording Scripts/Jordan_Recording_Script.docx','Jordan Recording Guide','Jordan','Student voice associated with the Confident Student Problem; repository paths retain legacy scenario numbering.','Review current displayed numbering and dialogue before recording or renaming any file.')
make_guide(DOCS/'Recording Scripts/Eli_Recording_Script.docx','Eli Recording Guide','Eli','Earlier Scenario 1 student voice retained in legacy assets and documentation. Maya is the current Scenario 1 student.','Do not record the old Content Avalanche line list. Keep Eli on hold unless a current scenario explicitly restores the character.')
print('wrote',len(images),'visuals',len(audio),'audio')

